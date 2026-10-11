from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
C_NAVY = RGBColor(15, 41, 66)
C_BLUE = RGBColor(0, 113, 206)
C_YELLOW = RGBColor(255, 194, 32)
C_TEAL = RGBColor(13, 148, 136)
C_CORAL = RGBColor(225, 29, 72)
C_SLATE = RGBColor(71, 85, 105)
C_DARK = RGBColor(30, 41, 59)
C_LIGHT = RGBColor(248, 250, 252)
C_BORDER = RGBColor(203, 213, 225)
C_WHITE = RGBColor(255, 255, 255)

def add_header(slide, title_text, subtitle_text, slide_num):
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_LIGHT
    bg.line.fill.background()

    # Top Navy Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05))
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_NAVY
    banner.line.fill.background()

    # Yellow Accent Strip
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.05), Inches(13.333), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = C_YELLOW
    accent.line.fill.background()

    # Title & Subtitle
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(11.5), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = 'Calibri'
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(186, 230, 253)

    # Footer
    ft = slide.shapes.add_textbox(Inches(0.5), Inches(7.12), Inches(12.333), Inches(0.3))
    ft_tf = ft.text_frame
    fp = ft_tf.paragraphs[0]
    fp.text = f"Dev Dilip Dhapodkar (PRN: 24070521505)  |  Predictive Analytics CA3 — Topic: Sales Forecasting                                                           Slide {slide_num} / 13"
    fp.font.name = 'Calibri'
    fp.font.size = Pt(9.5)
    fp.font.color.rgb = C_SLATE

def add_card(slide, left, top, width, height, title, bullets, header_color=C_BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.adjustments[0] = 0.02
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_BORDER
    card.line.width = Pt(1)

    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.42))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = header_color
    hdr.line.fill.background()

    htb = slide.shapes.add_textbox(Inches(left + 0.12), Inches(top + 0.04), Inches(width - 0.24), Inches(0.35))
    hp = htb.text_frame.paragraphs[0]
    hp.text = title
    hp.font.name = 'Calibri'
    hp.font.size = Pt(13)
    hp.font.bold = True
    hp.font.color.rgb = C_WHITE

    btb = slide.shapes.add_textbox(Inches(left + 0.14), Inches(top + 0.48), Inches(width - 0.28), Inches(height - 0.55))
    btf = btb.text_frame
    btf.word_wrap = True
    for i, b in enumerate(bullets):
        bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        bp.text = b
        bp.font.name = 'Calibri'
        bp.font.size = Pt(11.5)
        bp.font.color.rgb = C_DARK
        bp.space_after = Pt(5)

def add_kpi_box(slide, left, top, width, height, metric_val, metric_lbl, sub_lbl, accent_color=C_BLUE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.adjustments[0] = 0.04
    box.fill.solid()
    box.fill.fore_color.rgb = C_WHITE
    box.line.color.rgb = accent_color
    box.line.width = Pt(1.8)

    tb = slide.shapes.add_textbox(Inches(left + 0.1), Inches(top + 0.08), Inches(width - 0.2), Inches(height - 0.15))
    tf = tb.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.text = metric_val
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = 'Calibri'
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = accent_color

    p2 = tf.add_paragraph()
    p2.text = metric_lbl
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = 'Calibri'
    p2.font.size = Pt(11.5)
    p2.font.bold = True
    p2.font.color.rgb = C_NAVY

    p3 = tf.add_paragraph()
    p3.text = sub_lbl
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = 'Calibri'
    p3.font.size = Pt(9.5)
    p3.font.color.rgb = C_SLATE

def add_notes(slide, text):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text

def add_styled_table(slide, left, top, width, height, rows_data, col_widths=None):
    R = len(rows_data)
    C = len(rows_data[0])
    shape = slide.shapes.add_table(R, C, Inches(left), Inches(top), Inches(width), Inches(height))
    table = shape.table
    if col_widths:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Inches(w)
    for r in range(R):
        for c in range(C):
            cell = table.cell(r, c)
            cell.text = str(rows_data[r][c])
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = C_NAVY
            elif r % 2 == 1:
                cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
            else:
                cell.fill.fore_color.rgb = C_WHITE
            for p in cell.text_frame.paragraphs:
                p.font.name = 'Calibri'
                p.font.size = Pt(10.5 if r > 0 else 11)
                p.font.bold = (r == 0 or c == 0)
                p.font.color.rgb = C_WHITE if r == 0 else C_DARK

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY
bg1.line.fill.background()

strip1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.28), Inches(7.5))
strip1.fill.solid()
strip1.fill.fore_color.rgb = C_YELLOW
strip1.line.fill.background()

tb1 = s1.shapes.add_textbox(Inches(0.85), Inches(0.85), Inches(11.8), Inches(3.6))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "PREDICTIVE ANALYTICS — CONTINUOUS ASSESSMENT 3 (CA3)"
p.font.name = 'Calibri'
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_YELLOW
p.space_after = Pt(12)

p = tf1.add_paragraph()
p.text = "Sales Forecasting:\nEnd-to-End Retail Store Sales & Supply Chain Optimization"
p.font.name = 'Calibri'
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = C_WHITE
p.space_after = Pt(14)

p = tf1.add_paragraph()
p.text = "Forecasting Weekly Department-Level Retail Demand Across 45 Walmart Stores Using Historical Sales, Promotional Markdowns, Holiday Events, and Macroeconomic Indicators via the KDD Process"
p.font.name = 'Calibri'
p.font.size = Pt(14.5)
p.font.color.rgb = RGBColor(186, 230, 253)
p.space_after = Pt(14)

p_auth = tf1.add_paragraph()
p_auth.text = "Presented By: Dev Dilip Dhapodkar   |   PRN: 24070521505"
p_auth.font.name = 'Calibri'
p_auth.font.size = Pt(16)
p_auth.font.bold = True
p_auth.font.color.rgb = C_YELLOW

# 4 Hero KPI Cards on Title Slide
add_kpi_box(s1, 0.85, 4.75, 2.75, 1.55, "421,570", "Training Records", "45 Stores × 81 Depts (143 Wks)", C_BLUE)
add_kpi_box(s1, 3.85, 4.75, 2.75, 1.55, "43 Features", "Leak-Free KDD Pipeline", "Lags, Rolling, Promo & Macro", C_TEAL)
add_kpi_box(s1, 6.85, 4.75, 2.75, 1.55, "$1,758.94", "Best WMAE (LightGBM)", "-60.0% Error vs. Baseline", C_CORAL)
add_kpi_box(s1, 9.85, 4.75, 2.75, 1.55, "R² = 0.9725", "Variance Explained", "$31.7M (-39.5%) Cost Savings", C_BLUE)

add_notes(s1, "Good morning/afternoon. Today I am presenting my End-to-End Predictive Model Development Project for Predictive Analytics CA3: Retail Store Sales Forecasting and Supply Chain Inventory Optimization on the official Kaggle Walmart dataset following the KDD process.")

# ==============================================================================
# SLIDE 2: PROBLEM STATEMENT & SUPPLY CHAIN MOTIVATION
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "1. Problem Statement & Supply Chain Optimization Objective", "Balancing Out-of-Stock (OOS) Risk Against Inventory Holding Costs Across Multi-Format Retail Stores", 2)

add_card(s2, 0.5, 1.35, 6.0, 2.65, "The Retail Forecasting Challenge", [
    "• Granular Complexity: Forecasting Weekly_Sales across 45 stores and 81 distinct departments (3,645 series) over 143 weeks.",
    "• Exogenous Volatility: Demand is simultaneously driven by 5 promotional markdown channels (MarkDown1–5), 4 major holiday events, and regional macroeconomic shifts (CPI, Fuel Price, Unemployment).",
    "• Format Heterogeneity: Stores span Type A Supercenters (22), Type B Discount Stores (17), and Type C Neighborhood Markets (6)."
], C_NAVY)

add_card(s2, 6.8, 1.35, 6.0, 2.65, "Asymmetric Supply Chain Cost Structure", [
    "• Under-Forecasting (Out-of-Stock / OOS): Stockouts during high-velocity holiday peaks cause immediate lost revenue and customer attrition.",
    "• 5x Holiday Weighting (WMAE): Official competition metric assigns weight w = 5 to Holiday weeks and w = 1 to Regular weeks.",
    "• Over-Forecasting (Holding Cost): Excess safety stock inflates warehouse storage overhead, working capital tie-up, and clearance markdowns."
], C_CORAL)

add_card(s2, 0.5, 4.25, 12.3, 2.65, "Primary Project Objectives (KDD Framework)", [
    "1. Relational Ingestion & Preprocessing: Merge train (421,570 rows), features (8,190 rows), and stores (45 rows); impute structural missingness in MarkDown1–5 (64%–74%) and treat negative accounting returns (1,285 rows) & extreme outliers.",
    "2. Multi-Horizon Feature Engineering: Engineer 43 leak-free features including seasonal lags (Lag_1, Lag_2, Lag_4, Lag_52), t-1 shifted rolling averages (4-week & 12-week), Fourier cyclical encodings, and holiday-promo interaction terms.",
    "3. Multi-Model Benchmarking & Business Translation: Compare Baseline (Ridge Regression), Random Forest, XGBoost, and LightGBM across WMAE, RMSE, MAE, and R², and quantify inventory holding & stockout cost reductions."
], C_BLUE)

add_notes(s2, "In retail supply chains, forecast errors have asymmetric costs. Under-forecasting during holidays leads to severe stockouts, which is why the Kaggle WMAE metric weights holiday weeks 5 times higher than regular weeks. Over-forecasting increases holding costs.")

# ==============================================================================
# SLIDE 3: LITERATURE REVIEW & RESEARCH GAP
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "2. Literature Review & Research Gap", "Evolution from Classical Univariate Time Series to Global Gradient-Boosted Decision Trees", 3)

lit_rows = [
    ["Author(s) & Year", "Methodology", "Domain / Dataset", "Key Findings & Limitations"],
    ["Box & Jenkins (2015);\nHyndman (2021)", "ARIMA, SARIMAX &\nETS State-Space", "Univariate Retail\nTime Series", "Captures linear autocorrelation; fails on abrupt non-linear promotional spikes and cross-store hierarchies."],
    ["Chen & Guestrin\n(2016)", "XGBoost (Regularized\nGradient Boosting)", "Large-Scale Tabular &\nTemporal Benchmarks", "2nd-order Taylor objective with L1/L2 tree penalties handles sparse markdowns and non-linear interactions."],
    ["Ke et al. (2017)", "LightGBM (Histogram\nLeaf-Wise Boosting)", "High-Dimensional\nMulti-Store Retail", "Leaf-wise splitting & histogram binning cut training time by >80% while boosting accuracy on skewed targets."],
    ["Makridakis et al.\n(2022)", "M5 Forecasting\nCompetition", "Walmart Hierarchical\nSales (42,840 Series)", "Global gradient boosted trees with multi-horizon lags & rolling windows outperformed pure ARIMA & deep nets."],
    ["Barrow & Kourentzes\n(2016)", "Inventory Safety Stock\nOptimization", "Supply Chain\nReplenishment", "Reducing peak-season MAE directly lowers required safety stock buffers and stockout penalty exposure."]
]
add_styled_table(s3, 0.5, 1.35, 12.333, 3.85, lit_rows, [2.2, 2.2, 2.2, 5.733])

add_card(s3, 0.5, 5.4, 12.333, 1.55, "Research Gap Addressed in Our Implementation", [
    "• Prevention of Temporal Data Leakage: Many existing implementations compute rolling averages without shifting by t-1 or compute target encodings across the full dataset. We enforce strict chronological isolation (split at 2011-11-25) and t-1 shifted windows.",
    "• Asymmetric Holiday Cost Optimization: We integrate w = 5 sample weights directly into ensemble training and translate WMAE gains into dollarized supply chain holding and stockout cost savings across Store Types A, B, and C."
], C_TEAL)

add_notes(s3, "Our literature review highlights why global tree ensembles like XGBoost and LightGBM surpassed classical ARIMA in the M5 Walmart competition. We address two common gaps: eliminating rolling-window data leakage and directly optimizing for 5x holiday weights.")

# ==============================================================================
# SLIDE 4: END-TO-END KDD WORKFLOW ARCHITECTURE
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "3. End-to-End KDD Implementation Workflow", "Systematic Seven-Stage Knowledge Discovery in Databases (KDD) Pipeline", 4)

stages = [
    ("Stage 1: Data Selection", "Ingest 4 relational Kaggle CSVs (train: 421,570 rows, features: 8,190, stores: 45, test: 115,064) & merge on (Store, Date, IsHoliday).", C_NAVY),
    ("Stage 2: Preprocessing", "Extract Has_MarkDown flags; zero-impute MarkDown1–5 (64–74% NA); ffill/bfill CPI & Unemp; floor 1,285 neg. sales & cap 99.95% outliers.", C_BLUE),
    ("Stage 3: Feature Eng.", "Construct 43 leak-free features: Lags (1, 2, 4, 52w), t-1 Rolling Means/Std (4w, 12w), Fourier seasonality, Holiday & Promo flags.", C_TEAL),
    ("Stage 4: Exploratory EDA", "Analyze store format hierarchy (Types A, B, C), physical size elasticity (r=0.81), holiday revenue lift (+44%), and macro correlations.", C_BLUE),
    ("Stage 5: Model Training", "Chronological split (263,199 train / 145,237 val rows); train Ridge Baseline, Random Forest, LightGBM, and XGBoost with w=5 weights.", C_NAVY),
    ("Stage 6: Evaluation", "Evaluate out-of-sample generalization via official Kaggle WMAE (w=5 holiday, w=1 regular), RMSE, MAE, R², and Holiday MAE.", C_CORAL),
    ("Stage 7: Interpretation", "Analyze XGBoost/LightGBM split gain attribution and quantify $31.7M (-39.5%) supply chain inventory mismatch cost reduction.", C_TEAL)
]

for idx, (st_name, st_desc, st_col) in enumerate(stages):
    left_pos = 0.45 + idx * 1.8
    box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.45), Inches(1.68), Inches(3.25))
    box.adjustments[0] = 0.04
    box.fill.solid()
    box.fill.fore_color.rgb = C_WHITE
    box.line.color.rgb = st_col
    box.line.width = Pt(2)

    hdr = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_pos), Inches(1.45), Inches(1.68), Inches(0.55))
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = st_col
    hdr.line.fill.background()

    htb = s4.shapes.add_textbox(Inches(left_pos + 0.05), Inches(1.48), Inches(1.58), Inches(0.5))
    hp = htb.text_frame.paragraphs[0]
    hp.text = st_name
    hp.alignment = PP_ALIGN.CENTER
    hp.font.name = 'Calibri'
    hp.font.size = Pt(11)
    hp.font.bold = True
    hp.font.color.rgb = C_WHITE

    dtb = s4.shapes.add_textbox(Inches(left_pos + 0.08), Inches(2.08), Inches(1.52), Inches(2.5))
    dtf = dtb.text_frame
    dtf.word_wrap = True
    dp = dtf.paragraphs[0]
    dp.text = st_desc
    dp.font.name = 'Calibri'
    dp.font.size = Pt(10.5)
    dp.font.color.rgb = C_DARK

add_card(s4, 0.45, 4.95, 12.42, 1.95, "Strict Chronological Validation Design (Preventing Look-Ahead Bias)", [
    "• Training Window (94 Weeks): 2010-02-05 to 2011-11-18 (263,199 store-department-week observations after lag initialization).",
    "• Validation Window (49 Weeks): 2011-11-25 to 2012-10-26 (145,237 observations) — spans a full seasonal cycle covering ALL FOUR official holiday events: Thanksgiving (Nov 2011), Christmas (Dec 2011), Super Bowl (Feb 2012), and Labor Day (Sep 2012)."
], C_NAVY)

add_notes(s4, "This slide illustrates our 7-stage KDD workflow from Data Selection and Preprocessing through Feature Engineering, EDA, Model Training, WMAE Evaluation, and Supply Chain Cost Interpretation.")

# ==============================================================================
# SLIDE 5: DATA SELECTION & PREPROCESSING (MISSING MARKDOWN & OUTLIERS)
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "4. Data Selection, Missing MarkDown Imputation & Outlier Treatment", "Relational Schema Integration and Domain-Driven Preprocessing of 421,570 Historical Sales Records", 5)

s5.shapes.add_picture("figures/fig2_outlier_and_markdown_analysis.png", Inches(0.45), Inches(1.3), width=Inches(12.43))

prep_rows = [
    ["Data Quality Issue", "Affected Records", "Root Cause in Retail Operations", "KDD Preprocessing & Treatment"],
    ["MarkDown1–MarkDown5 Missingness", "64.08% – 73.61% in Train", "Walmart began tracking promo markdowns in Nov 2011; NA = no active promo", "Extract Has_MarkDown & Active_MarkDown_Count; impute structural 0.0"],
    ["CPI & Unemployment Missingness", "585 rows (7.14%) in features.csv", "Macroeconomic series lagged in final 2013 test horizon weeks", "Store-grouped chronological forward/backward fill (ffill / bfill)"],
    ["Negative Weekly_Sales (< $0.00)", "1,285 rows (0.30%) in Train", "Post-holiday customer returns and accounting credit adjustments", "Floor negative sales at $0.00 to ensure non-negative demand target"],
    ["Extreme Upper-Tail Outliers", "422 rows > 99.9th pct ($174.9K)", "One-off bulk institutional clearances vs. genuine holiday surges", "Winsorize non-holiday records > 99.95th pct ($183.5K); preserve Holidays"]
]
add_styled_table(s5, 0.45, 4.95, 12.43, 2.0, prep_rows, [2.6, 2.1, 3.8, 3.93])

add_notes(s5, "In Step 2 of our workflow, we handle two crucial preprocessing tasks: first, MarkDown1 to MarkDown5 have 64% to 74% missing values prior to November 2011. We extract binary promo flags before zero-imputation. Second, we floor 1,285 negative sales rows at zero and winsorize non-holiday outliers.")

# ==============================================================================
# SLIDE 6: EXPLORATORY DATA ANALYSIS — STORE FORMATS & SEASONALITY
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "5. Exploratory Data Analysis: Store Formats & Seasonal Surges", "Multi-Store Sales Trajectory (2010–2012), Format Hierarchy (Types A, B, C), and Size Elasticity", 6)

s6.shapes.add_picture("figures/fig1_eda_sales_overview.png", Inches(0.45), Inches(1.25), width=Inches(8.55))

add_card(s6, 9.15, 1.25, 3.7, 5.65, "Key Visual EDA Insights", [
    "• Q4 Holiday Amplification: Aggregate weekly sales across the 45 stores average $47.1M in regular weeks, surging to $66.6M on Thanksgiving and $79.8M in the pre-Christmas week.",
    "• Store Format Hierarchy:\n  - Type A (22 Supercenters): Mean $1.45M/week (Avg Size: 177K sq.ft.)\n  - Type B (17 Discount Stores): Mean $0.79M/week (Avg Size: 101K sq.ft.)\n  - Type C (6 Neighborhood Mkts): Mean $0.46M/week (Avg Size: 41K sq.ft.)",
    "• Strong Size Elasticity: Physical store area correlates strongly with weekly revenue (Pearson r = 0.81), motivating Store_Size and Store_Type_Code features.",
    "• Right-Skewed Dept Sales: Department sales span from $0 to >$600K with median $7,612 and IQR fence at $47,395."
], C_NAVY)

add_notes(s6, "Figure 1 shows the aggregate weekly sales across all 45 Walmart stores from February 2010 to October 2012. Notice the massive recurring Q4 spikes during Thanksgiving and pre-Christmas weeks, and the strong r = 0.81 correlation between physical store square footage and revenue.")

# ==============================================================================
# SLIDE 7: EDA — HOLIDAY EVENTS, MACROECONOMICS & CORRELATIONS
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "6. EDA: Holiday Revenue Lift, Macroeconomics & Correlations", "Quantifying the 4 Official Holiday Events, Macroeconomic Trends, and Autoregressive Feature Correlations", 7)

s7.shapes.add_picture("figures/fig3_holiday_and_macro_effects.png", Inches(0.45), Inches(1.28), width=Inches(7.35))
s7.shapes.add_picture("figures/fig4_correlation_matrix.png", Inches(7.95), Inches(1.28), width=Inches(4.9))

add_card(s7, 0.45, 4.55, 7.35, 2.35, "Holiday Calendar & Macroeconomic Findings", [
    "• Thanksgiving Surge (+44.0%): Store revenue jumps from $1.02M (regular week) to $1.47M during Thanksgiving (Black Friday week) and $1.43M (+40.2%) during Pre-Christmas Weeks 49–51.",
    "• Super Bowl (+5.6%) & Labor Day (+2.0%): Generate moderate category-specific surges in grocery, electronics, and apparel departments, whereas Christmas Week drops (-6.0%) due to store closures.",
    "• Macroeconomic Context: Over 2010–2012, CPI rose steadily (+7.2%) while Unemployment declined from 8.6% to 7.1%, providing regional purchasing-power context across store markets."
], C_BLUE)

add_notes(s7, "Figure 3 quantifies the exact revenue lift of each holiday event: +44% during Thanksgiving and +40.2% during pre-Christmas weeks. Figure 4 shows the correlation matrix, confirming that Lag_52 (r=0.97), Rolling_Mean_4 (r=0.96), and Lag_1 (r=0.95) are strongly correlated with Weekly_Sales.")

# ==============================================================================
# SLIDE 8: FEATURE ENGINEERING ARCHITECTURE (43 PREDICTORS)
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "7. Feature Engineering Pipeline (43 Domain-Specific Predictors)", "Capturing Multi-Horizon Autoregression, Rolling Momentum, Cyclical Seasonality & Promo Interactions", 8)

add_card(s8, 0.45, 1.35, 4.0, 2.7, "1. Autoregressive Lag Features", [
    "• Grouped strictly by (Store, Dept) and sorted chronologically by Date:",
    "• Lag_1 (t - 1 week): Captures immediate local demand level (r = 0.95).",
    "• Lag_2 (t - 2 weeks): Captures bi-weekly paycheck cycles.",
    "• Lag_4 (t - 4 weeks): Captures monthly replenishment baseline (r = 0.94).",
    "• Lag_52 (t - 52 weeks): Exact year-over-year seasonal anchor (r = 0.97)."
], C_NAVY)

add_card(s8, 4.65, 1.35, 4.0, 2.7, "2. Leak-Free Rolling Windows", [
    "• Computed on t-1 shifted series s.shift(1).rolling(W) to prevent target leakage:",
    "• Rolling_Mean_4: 4-week moving average capturing short-term trend (r = 0.96).",
    "• Rolling_Std_4: 4-week rolling standard deviation measuring demand volatility.",
    "• Rolling_Mean_12: 12-week quarterly demand baseline (r = 0.95).",
    "• Momentum_1_4: Lag_1 - Rolling_Mean_4 (short-term acceleration)."
], C_BLUE)

add_card(s8, 8.85, 1.35, 4.0, 2.7, "3. Holiday & Promo Indicators", [
    "• Explicit Event Flags: Is_SuperBowl (W6), Is_LaborDay (W36), Is_Thanksgiving (W47), Is_Christmas (W52).",
    "• Pre_Christmas_Peak: Binary indicator for high-velocity shopping Weeks 49–51.",
    "• Promotional Aggregates: Total_MarkDown, Active_MarkDown_Count, Has_MarkDown.",
    "• Promo_Holiday_Interaction: Captures multiplicative promo lift during holidays."
], C_CORAL)

add_card(s8, 0.45, 4.25, 6.1, 2.65, "4. Cyclical Fourier & Calendar Features", [
    "• Standard Calendar: Year, Month, ISO Week (1–52), Quarter, DayOfYear.",
    "• Trigonometric Fourier Transforms:\n  - Week_Sin = sin(2π · Week / 52), Week_Cos = cos(2π · Week / 52)\n  - Month_Sin = sin(2π · Month / 12), Month_Cos = cos(2π · Month / 12)",
    "• Preserves smooth circular continuity between Week 52 (late Dec) and Week 1 (early Jan) for tree & linear models."
], C_TEAL)

add_card(s8, 6.75, 4.25, 6.1, 2.65, "5. Out-of-Sample Store-Dept & Macro Features", [
    "• Leak-Free Target Statistics: Store_Dept_Mean, Store_Dept_Median, and Store_Dept_Std computed strictly on the training split (< 2011-11-25).",
    "• Structural Store Metadata: Store_Type_Code (A=3, B=2, C=1) and physical Size (sq. ft.).",
    "• Macroeconomic & Weather Context: Regional CPI, Unemployment Rate (%), Fuel_Price ($/gal), and Temperature (°F)."
], C_NAVY)

add_notes(s8, "In Step 3, we engineered 43 domain-specific features across 5 families: autoregressive lags, leak-free t-1 shifted rolling averages, explicit holiday and promotional interaction flags, Fourier cyclical transforms, and out-of-sample store-department encodings.")

# ==============================================================================
# SLIDE 9: MODEL TRAINING & EVALUATION (WMAE, RMSE, MAE, R2)
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "8. Model Training & Quantitative Evaluation (WMAE, RMSE, R²)", "Benchmarking Baseline Ridge, Random Forest, XGBoost, and LightGBM on 145,237 Validation Dept-Weeks", 9)

s9.shapes.add_picture("figures/fig5_model_comparison.png", Inches(0.45), Inches(1.22), width=Inches(12.43))

eval_rows = [
    ["Model Architecture", "Key Hyperparameters (Sample-Weighted w=5)", "WMAE ($ USD) ↓", "RMSE ($ USD) ↓", "MAE ($ USD) ↓", "R² Score ↑", "Holiday MAE ↓", "Train Time"],
    ["Baseline (Ridge)", "StandardScaler + L2 Ridge (alpha = 100.0)", "$4,393.17", "$6,616.57", "$2,514.98", "0.9183", "$10,108.52", "0.20 s"],
    ["Random Forest", "100 trees, max_depth=16, max_features=0.6", "$1,791.54", "$4,413.86", "$1,520.75", "0.9637", "$2,615.57", "42.88 s"],
    ["XGBoost Regressor", "550 trees, lr=0.05, max_depth=8, L1=1, L2=5", "$1,837.57", "$4,306.19", "$1,523.89", "0.9654", "$2,792.11", "6.31 s"],
    ["LightGBM (Best WMAE)", "600 trees, lr=0.05, num_leaves=63,colsample=0.8", "$1,758.94", "$3,837.62", "$1,541.31", "0.9725", "$2,421.19", "4.91 s"]
]
add_styled_table(s9, 0.45, 4.65, 12.43, 2.25, eval_rows, [2.1, 3.33, 1.25, 1.25, 1.15, 0.95, 1.3, 1.1])

add_notes(s9, "Here are our quantitative validation results across 145,237 department-weeks. LightGBM achieves the lowest WMAE of $1,758.94—a 60% reduction over the Baseline Ridge model—along with the best RMSE ($3,837.62) and highest R-squared (0.9725) in under 5 seconds.")

# ==============================================================================
# SLIDE 10: CHRONOLOGICAL FORECAST TRACKING & PARITY ANALYSIS
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "9. Validation Forecast Tracking & Parity Analysis", "Tracking 49 Chronological Validation Weeks (Nov 2011 – Oct 2012) Across All 4 Major Holidays", 10)

s10.shapes.add_picture("figures/fig6_actual_vs_predicted.png", Inches(0.45), Inches(1.25), width=Inches(12.43))

add_card(s10, 0.45, 5.25, 12.43, 1.7, "Why Tree Ensembles Outperform Linear Models During Holiday Peaks", [
    "• Linear Overshoot & Lag Bias (Panel A): Ridge Regression relies linearly on Lag_1 and Rolling_Mean_4, causing massive overshoot immediately after Thanksgiving/Christmas spikes (Holiday MAE = $10,108.52).",
    "• Non-Linear Holiday Switching: LightGBM ($1,758.94 WMAE) and XGBoost ($1,837.57 WMAE, $1,410.37 Regular MAE) condition on Is_Thanksgiving, Pre_Christmas_Peak, and Lag_52 simultaneously, tracking both the $77M Q4 peak and the post-Christmas drop with R² > 0.965."
], C_NAVY)

add_notes(s10, "Figure 6 Panel A shows the weekly forecast tracking across the 49-week validation period. While the linear baseline overshoots massively after holiday spikes, LightGBM and XGBoost track the Thanksgiving, Christmas, Super Bowl, and Labor Day surges tightly.")

# ==============================================================================
# SLIDE 11: FEATURE IMPORTANCE & GAIN ATTRIBUTION
# ==============================================================================
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "10. Feature Importance & Gain Attribution Analysis", "Comparing Normalized Split Gain Across XGBoost and LightGBM Regressors", 11)

s11.shapes.add_picture("figures/fig7_feature_importance.png", Inches(0.45), Inches(1.25), width=Inches(12.43))

add_card(s11, 0.45, 5.3, 12.43, 1.65, "Domain Interpretation of Top Predictive Drivers", [
    "• Annual Seasonal Anchor (Lag_52): Ranked #1 in both models (31.4% XGBoost gain, 40.9% LightGBM gain), proving that department-specific sales from the exact same week a year ago provide the strongest structural baseline.",
    "• Short-Term Momentum (Rolling_Mean_4 & Lag_1): Account for 37.6% (XGBoost) and 39.8% (LightGBM) of total gain, dynamically adapting forecasts to current store traffic, inflation (CPI), and local macroeconomic conditions.",
    "• Explicit Holiday Indicators (Is_Thanksgiving, Is_Christmas, Pre_Christmas_Peak): Contribute 9.4% gain in XGBoost, driving accurate Q4 surge adjustments."
], C_BLUE)

add_notes(s11, "Figure 7 compares feature importance across XGBoost and LightGBM. Year-over-year Lag_52 is the #1 predictor in both models (31.4% and 40.9%), followed by Rolling_Mean_4, Lag_1, Store_Dept_Mean, and explicit Thanksgiving and Christmas indicators.")

# ==============================================================================
# SLIDE 12: SUPPLY CHAIN INVENTORY OPTIMIZATION & COST IMPACT
# ==============================================================================
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "11. Supply Chain Inventory Optimization & Business Impact", "Translating Forecast Accuracy into Safety Stock & Stockout Cost Savings Across Store Types A, B, and C", 12)

s12.shapes.add_picture("figures/fig8_supply_chain_impact.png", Inches(0.45), Inches(1.25), width=Inches(12.43))

add_card(s12, 0.45, 5.15, 12.43, 1.8, "Operational Supply Chain Impact (49-Week Validation Horizon)", [
    "• 76.0% Reduction in Holiday Peak Error (Panel A): LightGBM cuts Holiday MAE from $10,108.52 to $2,421.19 per department-week, directly preventing high-margin out-of-stock (OOS) occurrences during Black Friday and Christmas.",
    "• $31.7M (-39.5%) Mismatch Cost Reduction (Panel B): Applying an industry-standard 22% holding + stockout penalty rate on cumulative forecast error reduces inventory mismatch cost from $80.3M (Baseline) to $48.6M:\n  - Type A Supercenters (22 stores): $46.5M → $29.4M (-36.7% / $17.1M saved)  |  Type B (17 stores): $29.7M → $16.9M (-43.0%)  |  Type C (6 stores): $4.1M → $2.3M (-43.5%)."
], C_TEAL)

add_notes(s12, "Figure 8 translates our statistical metrics into supply chain impact. Cutting holiday error by 76% prevents peak-season stockouts, while overall error reduction lowers estimated holding and stockout mismatch costs by $31.7 million (-39.5%) across the 45 stores.")

# ==============================================================================
# SLIDE 13: CONCLUSION, RECOMMENDATIONS & FUTURE SCOPE
# ==============================================================================
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "12. Conclusion, Operational Recommendations & Future Scope", "Summary of End-to-End KDD Achievements and Next-Generation Retail Forecasting Roadmap", 13)

add_card(s13, 0.45, 1.35, 6.0, 2.75, "Key KDD Project Achievements", [
    "• Complete End-to-End Pipeline: Ingested & merged 421,570 historical sales records across 45 stores and 81 departments (143 weeks).",
    "• Leak-Free Feature Engineering: Engineered 43 predictors (Lags 1–52, t-1 Rolling Means, Fourier seasonality, Holiday & Promo flags).",
    "• Superior Accuracy: LightGBM achieved WMAE = $1,758.94 (-60.0% vs. Baseline), RMSE = $3,837.62, and R² = 0.9725 in 4.91s.",
    "• Business Value: $31.7M (-39.5%) reduction in estimated supply chain holding & stockout mismatch penalties."
], C_NAVY)

add_card(s13, 6.8, 1.35, 6.0, 2.75, "Actionable Supply Chain Recommendations", [
    "1. Hybrid Model Deployment: Deploy LightGBM ($2,421 Holiday MAE) for Q4 holiday inventory allocation and XGBoost ($1,410 Regular MAE) for routine weekly automated store replenishment.",
    "2. Pre-Holiday Safety Stock Staging: Stage safety stock 3 weeks ahead of Weeks 47 (Thanksgiving) and 50 (Pre-Christmas) when MarkDown2 and MarkDown3 promotions are scheduled.",
    "3. Format-Specific Reorder Points: Calibrate safety stock buffers separately for Type A Supercenters vs. Type B/C stores."
], C_BLUE)

add_card(s13, 0.45, 4.35, 12.35, 2.5, "Future Scope & Advanced Extensions", [
    "• Probabilistic Quantile Forecasting: Extend LightGBM/XGBoost point predictions to quantile regression (Pinball Loss at τ = 0.50, 0.90, 0.95) and Conformal Prediction to directly compute service-level safety stockreorder points.",
    "• Hierarchical Forecast Reconciliation (MinT): Enforce bottom-up and top-down coherence across National → Store Format → Store → Department levels.",
    "• Deep Learning Architectures: Benchmark Temporal Fusion Transformers (TFT) and N-BEATS alongside gradient-boosted decision trees."
], C_TEAL)

add_notes(s13, "In conclusion, our KDD pipeline reduced WMAE by 60% to $1,758.94 and explained 97.25% of sales variance. We recommend a hybrid deployment using LightGBM for holiday peaks and XGBoost for regular replenishment. Thank you, and I welcome any questions.")

prs.save("Walmart_Sales_Forecasting_Presentation.pptx")
print("Saved Walmart_Sales_Forecasting_Presentation.pptx (13 slides) successfully!")
